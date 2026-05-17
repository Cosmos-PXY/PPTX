# -*- coding: utf-8 -*-
"""
将 input.png 转为一页可编辑 PPTX。
每一行 OCR 文本会被放到单独文本框，底图为原图。
生成文件：研究现状_单页复刻.pptx
依赖：python-pptx, pillow, pytesseract, opencv-python
需要系统安装 tesseract 和中文语言包 chi_sim

使用方法：
 1. 把你的图片保存为当前目录下的 input.png
 2. 在终端安装 tesseract（并安装 chi_sim），以及 Python 依赖
 3. 运行 python3 make_pptx_from_image.py

此脚本基于之前对话中提供的 OCR -> PPTX 方案，尽量把每行文本创建为独立文本框（字体：Microsoft YaHei）。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image
import pytesseract
import cv2
from pytesseract import Output

INPUT_IMAGE = "input.png"
OUTPUT_PPTX = "研究现状_单页复刻.pptx"
DPI = 96.0  # 假定屏幕 DPI 为 96，若你有不同 DPI 可修改


def px_to_inches(px, dpi=DPI):
    return px / dpi


def group_lines(ocr_data):
    n_boxes = len(ocr_data['level'])
    lines = {}
    for i in range(n_boxes):
        text = ocr_data['text'][i].strip()
        conf = None
        try:
            conf = int(float(ocr_data['conf'][i]))
        except Exception:
            conf = -1
        if text == "" or conf < 30:
            continue
        key = (ocr_data['block_num'][i], ocr_data['par_num'][i], ocr_data['line_num'][i])
        left = ocr_data['left'][i]
        top = ocr_data['top'][i]
        width = ocr_data['width'][i]
        height = ocr_data['height'][i]
        if key not in lines:
            lines[key] = {
                'texts': [text],
                'left': left,
                'top': top,
                'right': left + width,
                'bottom': top + height,
                'confs': [conf],
            }
        else:
            lines[key]['texts'].append(text)
            lines[key]['confs'].append(conf)
            lines[key]['left'] = min(lines[key]['left'], left)
            lines[key]['top'] = min(lines[key]['top'], top)
            lines[key]['right'] = max(lines[key]['right'], left + width)
            lines[key]['bottom'] = max(lines[key]['bottom'], top + height)
    results = []
    for k, v in lines.items():
        text_line = " ".join(v['texts'])
        left = v['left']
        top = v['top']
        width = v['right'] - v['left']
        height = v['bottom'] - v['top']
        avg_conf = sum(v['confs']) / len(v['confs'])
        results.append({
            'text': text_line,
            'left': left,
            'top': top,
            'width': width,
            'height': height,
            'conf': avg_conf
        })
    results.sort(key=lambda x: (x['top'], x['left']))
    return results


def guess_font_size(px_height, dpi=DPI):
    factor = 0.85
    pt = int(max(8, px_height / dpi * 72 * factor))
    return pt


def main():
    if not os.path.exists(INPUT_IMAGE):
        print(f"错误：找不到输入图片 {INPUT_IMAGE}。请把图片保存为当前目录下的 {INPUT_IMAGE}")
        return

    pil_img = Image.open(INPUT_IMAGE)
    img_w, img_h = pil_img.size

    prs = Presentation()
    prs.slide_width = Inches(img_w / DPI)
    prs.slide_height = Inches(img_h / DPI)

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    left = top = Inches(0)
    pic = slide.shapes.add_picture(INPUT_IMAGE, left, top, width=prs.slide_width, height=prs.slide_height)

    try:
        img_cv = cv2.cvtColor(cv2.imread(INPUT_IMAGE), cv2.COLOR_BGR2RGB)
    except Exception as e:
        img_cv = cv2.imread(INPUT_IMAGE)
    ocr_data = pytesseract.image_to_data(img_cv, output_type=Output.DICT, lang='chi_sim+eng')
    lines = group_lines(ocr_data)

    for item in lines:
        txt = item['text']
        if txt.strip() == "":
            continue
        left_px = item['left']
        top_px = item['top']
        w_px = max(10, item['width'])
        h_px = max(10, item['height'])

        left_in = px_to_inches(left_px)
        top_in = px_to_inches(top_px)
        w_in = px_to_inches(w_px)
        h_in = px_to_inches(h_px)

        txBox = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(w_in), Inches(h_in))
        tf = txBox.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = txt
        p.alignment = PP_ALIGN.LEFT

        run = p.runs[0]
        run.font.name = 'Microsoft YaHei'
        font_size = guess_font_size(h_px)
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    prs.save(OUTPUT_PPTX)
    print(f"已生成 PPT 文件：{OUTPUT_PPTX}。请用 PowerPoint / WPS 打开并检查/微调文本框。")

if __name__ == "__main__":
    main()
